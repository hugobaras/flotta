"""Lecture de tableaux CSV et Excel envoyés par l'utilisateur (flotte, commandes)."""
import csv
import io
import unicodedata

from fastapi import HTTPException

MAX_BYTES = 5 * 1024 * 1024
MAX_LIGNES = 5000


def normaliser(entete: str) -> str:
    """« Hauteur (m) » -> « hauteur_m » : minuscules, sans accents, mots séparés par _."""
    texte = unicodedata.normalize("NFKD", entete).encode("ascii", "ignore").decode().lower()
    return "_".join("".join(c if c.isalnum() else " " for c in texte).split())


def _lire_csv(data: bytes) -> list[list[str]]:
    try:
        texte = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        texte = data.decode("cp1252", errors="replace")  # export Excel français
    echantillon = texte[:4096]
    try:
        dialecte = csv.Sniffer().sniff(echantillon, delimiters=";,\t")
    except csv.Error:
        dialecte = csv.excel
    return [[c.strip() for c in ligne] for ligne in csv.reader(io.StringIO(texte), dialecte)]


def _lire_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        classeur = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(422, "Fichier Excel illisible")
    lignes = []
    for ligne in classeur.worksheets[0].iter_rows(values_only=True):
        lignes.append(["" if c is None else str(c).strip() for c in ligne])
        if len(lignes) > MAX_LIGNES + 1:
            break
    classeur.close()
    return lignes


def lire_tableau(data: bytes, nom_fichier: str) -> tuple[list[str], list[list[str]]]:
    """Renvoie (en-têtes, lignes) ; les lignes entièrement vides sont ignorées."""
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "Fichier trop lourd (5 Mo maximum)")
    if nom_fichier.lower().endswith(".xlsx") or data[:2] == b"PK":
        lignes = _lire_xlsx(data)
    else:
        lignes = _lire_csv(data)
    lignes = [l for l in lignes if any(c for c in l)]
    if len(lignes) < 2:
        raise HTTPException(422, "Le fichier doit contenir une ligne d'en-têtes et au moins une ligne de données")
    if len(lignes) - 1 > MAX_LIGNES:
        raise HTTPException(422, f"{MAX_LIGNES} lignes maximum par fichier")
    return lignes[0], lignes[1:]


def nombre(valeur: str) -> float | None:
    """Accepte la virgule décimale française ; vide -> None ; sinon ValueError."""
    v = valeur.strip().replace(" ", "").replace(" ", "").replace(",", ".")
    return float(v) if v else None


def booleen(valeur: str) -> bool:
    v = valeur.strip().lower()
    if v in ("", "non", "n", "0", "false", "faux"):
        return False
    if v in ("oui", "o", "1", "true", "vrai", "x"):
        return True
    raise ValueError("oui ou non attendu")
